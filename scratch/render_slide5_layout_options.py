import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import win32com.client
from PIL import Image, ImageDraw, ImageFont

# Define Color Palette
C_NAVY = RGBColor(15, 23, 42)
C_BLUE_TEMPLATE = RGBColor(37, 99, 235)
C_RED_TEMPLATE = RGBColor(220, 38, 38)
C_GREEN_TEMPLATE = RGBColor(22, 163, 74)
C_TEAL_TEMPLATE = RGBColor(13, 148, 136)
C_ORANGE_TEMPLATE = RGBColor(180, 83, 9)
C_TEXT_DARK = RGBColor(51, 65, 85)
C_TEXT_MUTED = RGBColor(100, 116, 139)
C_WHITE = RGBColor(255, 255, 255)
C_BLACK = RGBColor(0, 0, 0)

BASE_DIR = r"C:\Users\oshsh\.gemini\antigravity\scratch\sih26056_airfare_cpi"
ASSET_DIR = os.path.join(BASE_DIR, "public", "presentation")
OUTPUT_PPTX = os.path.join(BASE_DIR, "public", "presentation", "slide5_layout_options.pptx")

logo_path = os.path.join(ASSET_DIR, "sih_logo.png")
proto_chart = os.path.join(ASSET_DIR, "vayusuchak_prototype_chart.png")
qr_path = os.path.join(ASSET_DIR, "prototype_qr.png")

def add_clean_header(slide, option_pill_text=""):
    # Header container
    tb_title = slide.shapes.add_textbox(Inches(2.50), Inches(0.16), Inches(8.10), Inches(0.92))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0

    p_main = tf_title.paragraphs[0]
    p_main.text = "IMPACT AND BENEFITS"
    p_main.font.name = "Georgia"
    p_main.font.size = Pt(23)
    p_main.font.bold = True
    p_main.font.color.rgb = C_NAVY
    p_main.alignment = PP_ALIGN.CENTER

    p_sub = tf_title.add_paragraph()
    p_sub.text = "Data-Driven Macroeconomic Visibility & Inflation Accuracy"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(12)
    p_sub.font.bold = True
    p_sub.font.color.rgb = C_BLUE_TEMPLATE
    p_sub.alignment = PP_ALIGN.CENTER

    # Roorkies pill
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(0.22), Inches(1.80), Inches(0.68))
    pill.adjustments[0] = 0.5
    pill.fill.solid()
    pill.fill.fore_color.rgb = C_WHITE
    pill.line.color.rgb = C_NAVY
    pill.line.width = Pt(1.5)
    tf_pill = pill.text_frame
    p_pill = tf_pill.paragraphs[0]
    p_pill.text = "Roorkies"
    p_pill.font.name = "Arial"
    p_pill.font.size = Pt(13)
    p_pill.font.bold = True
    p_pill.font.color.rgb = C_BLACK
    p_pill.alignment = PP_ALIGN.CENTER

    # Option tag badge next to Roorkies pill
    if option_pill_text:
        opt_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.45), Inches(0.32), Inches(1.55), Inches(0.46))
        opt_box.adjustments[0] = 0.4
        opt_box.fill.solid()
        opt_box.fill.fore_color.rgb = RGBColor(238, 242, 255)
        opt_box.line.color.rgb = RGBColor(99, 102, 241)
        opt_box.line.width = Pt(1.2)
        tf_ob = opt_box.text_frame
        p_ob = tf_ob.paragraphs[0]
        p_ob.text = option_pill_text
        p_ob.font.name = "Arial"
        p_ob.font.size = Pt(9.5)
        p_ob.font.bold = True
        p_ob.font.color.rgb = RGBColor(67, 56, 202)
        p_ob.alignment = PP_ALIGN.CENTER

    # Logo
    if os.path.exists(logo_path):
        slide.shapes.add_picture(logo_path, Inches(10.75), Inches(0.12), Inches(1.85), Inches(1.05))

def add_bottom_banner(slide, y_pos=Inches(6.55), h_pos=Inches(0.70)):
    banner_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), y_pos, Inches(12.23), h_pos)
    banner_pill.adjustments[0] = 0.22
    banner_pill.fill.solid()
    banner_pill.fill.fore_color.rgb = C_TEAL_TEMPLATE
    banner_pill.line.color.rgb = RGBColor(15, 118, 110)
    banner_pill.line.width = Pt(1.5)

    tb_b = slide.shapes.add_textbox(Inches(0.72), y_pos + Inches(0.08), Inches(10.50), h_pos - Inches(0.16))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = 0

    p_b1 = tf_b.paragraphs[0]
    r_b1a = p_b1.add_run()
    r_b1a.text = "🚀 Live Interactive Prototype: "
    r_b1a.font.name = "Arial"
    r_b1a.font.size = Pt(12)
    r_b1a.font.bold = True
    r_b1a.font.color.rgb = C_WHITE

    r_b1b = p_b1.add_run()
    r_b1b.text = "https://sih26056-airfare-cpi.vercel.app ↗"
    r_b1b.font.name = "Arial"
    r_b1b.font.size = Pt(12)
    r_b1b.font.bold = True
    r_b1b.font.underline = True
    r_b1b.font.color.rgb = RGBColor(224, 242, 254)

    p_b2 = tf_b.add_paragraph()
    r_b2 = p_b2.add_run()
    r_b2.text = "Open API for ministries, regulators and researchers | Scan QR code on right to explore live"
    r_b2.font.name = "Arial"
    r_b2.font.size = Pt(11)
    r_b2.font.color.rgb = RGBColor(240, 253, 250)

    if os.path.exists(qr_path):
        slide.shapes.add_picture(qr_path, Inches(11.95), y_pos + Inches(0.06), Inches(0.68), Inches(0.58))

# Data items
impacts_list = [
    ("• Zero policy lag: ", "price collection cut from ~15 days to under 24 hours"),
    ("• Denser coverage: ", "3,650+ quotes/day across 12 routes x 4 airlines x 5 horizons, vs 1 quote per route per month"),
    ("• Captures what manual surveys miss: ", "lead-time effects and fare surges"),
    ("• Validated: ", "back-tested against official CPI air-fare component (Pearson r = 0.89; full sovereign series back-test planned with MoSPI NSO)")
]

benefit_boxes = [
    (
        "👥 SOCIAL",
        RGBColor(37, 99, 235), RGBColor(239, 246, 255),
        ("• Public measurement: ", "transparent, publicly checkable airfare-inflation data, relevant to UDAN regional travellers")
    ),
    (
        "📈 ECONOMIC",
        RGBColor(22, 163, 74), RGBColor(240, 253, 244),
        ("• Monetary policy signals: ", "leading transport-inflation indicator for RBI MPC; less manual field collection (savings under validation)")
    ),
    (
        "🌱 ENVIRONMENTAL",
        RGBColor(13, 148, 136), RGBColor(240, 253, 250),
        ("• Reduced field commutes: ", "less physical surveyor travel (qualitative)")
    ),
    (
        "🏛️ POLICY & INSTITUTIONAL",
        RGBColor(180, 83, 9), RGBColor(254, 252, 232),
        ("• Sovereign NSO feed: ", "automated API ingestion into MoSPI eSankhyiki and NDAP; abnormal fare surges flagged to DGCA for review")
    )
]

def build_option_1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_WHITE
    bg.line.fill.background()

    add_clean_header(s, "OPTION 1")

    # Left Column
    col_l_x = Inches(0.55)
    l_w = Inches(5.30)
    imp_y = Inches(1.15)
    imp_h = Inches(2.15)
    
    b_imp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col_l_x, imp_y, l_w, imp_h)
    b_imp.adjustments[0] = 0.04
    b_imp.fill.solid()
    b_imp.fill.fore_color.rgb = C_WHITE
    b_imp.line.color.rgb = C_BLUE_TEMPLATE
    b_imp.line.width = Pt(1.5)

    tb_imp = s.shapes.add_textbox(col_l_x + Inches(0.14), imp_y + Inches(0.08), l_w - Inches(0.28), imp_h - Inches(0.14))
    tf_imp = tb_imp.text_frame
    tf_imp.word_wrap = True
    tf_imp.margin_left = tf_imp.margin_top = tf_imp.margin_right = tf_imp.margin_bottom = 0

    p_ih = tf_imp.paragraphs[0]
    p_ih.text = "POTENTIAL IMPACT"
    p_ih.font.name = "Arial"
    p_ih.font.size = Pt(13)
    p_ih.font.bold = True
    p_ih.font.color.rgb = C_BLUE_TEMPLATE

    p_isub = tf_imp.add_paragraph()
    p_isub.text = "Target audience: MoSPI, RBI, DGCA, travellers"
    p_isub.font.name = "Arial"
    p_isub.font.size = Pt(9.5)
    p_isub.font.bold = True
    p_isub.font.color.rgb = C_TEXT_MUTED
    p_isub.space_after = Pt(4)

    for itit, idesc in impacts_list:
        p_i = tf_imp.add_paragraph()
        p_i.space_after = Pt(2.5)
        p_i.line_spacing = 1.08
        r_it = p_i.add_run()
        r_it.text = itit
        r_it.font.name = "Arial"
        r_it.font.size = Pt(10.5)
        r_it.font.bold = True
        r_it.font.color.rgb = C_BLACK
        r_id = p_i.add_run()
        r_id.text = idesc
        r_id.font.name = "Arial"
        r_id.font.size = Pt(10.5)
        r_id.font.color.rgb = C_TEXT_DARK

    ben_start_y = imp_y + imp_h + Inches(0.08)
    ben_box_h = Inches(0.68)
    ben_step = Inches(0.74)

    for b_idx, (b_title, b_theme, b_bg, (b_lbl, b_desc)) in enumerate(benefit_boxes):
        by = ben_start_y + b_idx * ben_step
        b_card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col_l_x, by, l_w, ben_box_h)
        b_card.adjustments[0] = 0.12
        b_card.fill.solid()
        b_card.fill.fore_color.rgb = b_bg
        b_card.line.color.rgb = b_theme
        b_card.line.width = Pt(1.2)

        tb_bc = s.shapes.add_textbox(col_l_x + Inches(0.12), by + Inches(0.03), l_w - Inches(0.24), ben_box_h - Inches(0.06))
        tf_bc = tb_bc.text_frame
        tf_bc.word_wrap = True
        tf_bc.margin_left = tf_bc.margin_top = tf_bc.margin_right = tf_bc.margin_bottom = 0

        p_bch = tf_bc.paragraphs[0]
        p_bch.text = b_title
        p_bch.font.name = "Arial"
        p_bch.font.size = Pt(11)
        p_bch.font.bold = True
        p_bch.font.color.rgb = b_theme
        p_bch.space_after = Pt(1)

        p_bl = tf_bc.add_paragraph()
        p_bl.line_spacing = 1.05
        r_bl1 = p_bl.add_run()
        r_bl1.text = b_lbl
        r_bl1.font.name = "Arial"
        r_bl1.font.size = Pt(10.5)
        r_bl1.font.bold = True
        r_bl1.font.color.rgb = C_BLACK
        r_bl2 = p_bl.add_run()
        r_bl2.text = b_desc
        r_bl2.font.name = "Arial"
        r_bl2.font.size = Pt(10.5)
        r_bl2.font.color.rgb = C_TEXT_DARK

    # Right Column
    r_x = Inches(6.05)
    r_w = Inches(6.73)
    r_h = Inches(5.26)
    box_p1 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, r_x, Inches(1.15), r_w, r_h)
    box_p1.adjustments[0] = 0.03
    box_p1.fill.solid()
    box_p1.fill.fore_color.rgb = C_WHITE
    box_p1.line.color.rgb = C_BLUE_TEMPLATE
    box_p1.line.width = Pt(1.5)

    tb_p1 = s.shapes.add_textbox(r_x + Inches(0.12), Inches(1.18), r_w - Inches(0.24), Inches(0.28))
    tf_p1 = tb_p1.text_frame
    p_p1 = tf_p1.paragraphs[0]
    r_p1a = p_p1.add_run()
    r_p1a.text = "PROTOTYPE IMAGE 1: Airfare Price Index Dashboard (Prototype data: sample)"
    r_p1a.font.name = "Arial"
    r_p1a.font.size = Pt(11)
    r_p1a.font.bold = True
    r_p1a.font.color.rgb = C_RED_TEMPLATE

    chart_y = Inches(1.48)
    chart_w = r_w - Inches(0.24)
    chart_h = Inches(4.35)
    if os.path.exists(proto_chart):
        s.shapes.add_picture(proto_chart, r_x + Inches(0.12), chart_y, chart_w, chart_h)

    tb_cap = s.shapes.add_textbox(r_x + Inches(0.12), Inches(5.87), chart_w, Inches(0.48))
    tf_cap = tb_cap.text_frame
    tf_cap.word_wrap = True
    p_cap = tf_cap.paragraphs[0]
    p_cap.alignment = PP_ALIGN.CENTER
    r_cap = p_cap.add_run()
    r_cap.text = "Geometric-mean (Jevons) aggregation avoids the upward bias of arithmetic averaging (Diewert, 2004). Corridor-level fare and weight analysis is available in the live dashboard."
    r_cap.font.name = "Arial"
    r_cap.font.size = Pt(10.5)
    r_cap.font.bold = True
    r_cap.font.color.rgb = C_NAVY

    add_bottom_banner(s)


def build_option_2(prs):
    # Option 2: Top 4-Column Benefits Grid (Full Width) + Bottom Split
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_WHITE
    bg.line.fill.background()

    add_clean_header(s, "OPTION 2")

    # Top 4 Columns (Full Width = 12.23 in)
    top_y = Inches(1.15)
    top_h = Inches(1.50)
    col_w = Inches(2.92)
    gap_x = Inches(0.18)
    start_x = Inches(0.55)

    for b_idx, (b_title, b_theme, b_bg, (b_lbl, b_desc)) in enumerate(benefit_boxes):
        cx = start_x + b_idx * (col_w + gap_x)
        b_card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, top_y, col_w, top_h)
        b_card.adjustments[0] = 0.08
        b_card.fill.solid()
        b_card.fill.fore_color.rgb = b_bg
        b_card.line.color.rgb = b_theme
        b_card.line.width = Pt(1.5)

        h_strip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, top_y, col_w, Inches(0.38))
        h_strip.adjustments[0] = 0.20
        h_strip.fill.solid()
        h_strip.fill.fore_color.rgb = b_theme
        h_strip.line.fill.background()
        tf_h = h_strip.text_frame
        p_h = tf_h.paragraphs[0]
        p_h.text = b_title
        p_h.font.name = "Arial"
        p_h.font.size = Pt(11)
        p_h.font.bold = True
        p_h.font.color.rgb = C_WHITE
        p_h.alignment = PP_ALIGN.CENTER

        tb_c = s.shapes.add_textbox(cx + Inches(0.10), top_y + Inches(0.42), col_w - Inches(0.20), top_h - Inches(0.46))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0

        p_c = tf_c.paragraphs[0]
        p_c.line_spacing = 1.08
        r_c1 = p_c.add_run()
        r_c1.text = b_lbl.replace("• ", "")
        r_c1.font.name = "Arial"
        r_c1.font.size = Pt(10)
        r_c1.font.bold = True
        r_c1.font.color.rgb = C_NAVY

        r_c2 = p_c.add_run()
        r_c2.text = b_desc
        r_c2.font.name = "Arial"
        r_c2.font.size = Pt(9.5)
        r_c2.font.color.rgb = C_TEXT_DARK

    # Bottom Split
    bot_y = Inches(2.78)
    bot_h = Inches(3.64)

    # Bottom Left: POTENTIAL IMPACT
    bot_l_w = Inches(5.30)
    b_imp2 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, start_x, bot_y, bot_l_w, bot_h)
    b_imp2.adjustments[0] = 0.04
    b_imp2.fill.solid()
    b_imp2.fill.fore_color.rgb = C_WHITE
    b_imp2.line.color.rgb = C_BLUE_TEMPLATE
    b_imp2.line.width = Pt(1.5)

    tb_imp2 = s.shapes.add_textbox(start_x + Inches(0.14), bot_y + Inches(0.12), bot_l_w - Inches(0.28), bot_h - Inches(0.24))
    tf_imp2 = tb_imp2.text_frame
    tf_imp2.word_wrap = True
    tf_imp2.margin_left = tf_imp2.margin_top = tf_imp2.margin_right = tf_imp2.margin_bottom = 0

    p_ih2 = tf_imp2.paragraphs[0]
    p_ih2.text = "🎯 POTENTIAL IMPACT"
    p_ih2.font.name = "Arial"
    p_ih2.font.size = Pt(13)
    p_ih2.font.bold = True
    p_ih2.font.color.rgb = C_BLUE_TEMPLATE

    p_isub2 = tf_imp2.add_paragraph()
    p_isub2.text = "Target audience: MoSPI, RBI, DGCA, travellers"
    p_isub2.font.name = "Arial"
    p_isub2.font.size = Pt(10)
    p_isub2.font.bold = True
    p_isub2.font.color.rgb = C_TEXT_MUTED
    p_isub2.space_after = Pt(6)

    for itit, idesc in impacts_list:
        p_i = tf_imp2.add_paragraph()
        p_i.space_after = Pt(5)
        p_i.line_spacing = 1.10
        r_it = p_i.add_run()
        r_it.text = itit
        r_it.font.name = "Arial"
        r_it.font.size = Pt(10.5)
        r_it.font.bold = True
        r_it.font.color.rgb = C_BLACK
        r_id = p_i.add_run()
        r_id.text = idesc
        r_id.font.name = "Arial"
        r_id.font.size = Pt(10)
        r_id.font.color.rgb = C_TEXT_DARK

    # Bottom Right: PROTOTYPE IMAGE 1
    bot_r_x = Inches(6.03)
    bot_r_w = Inches(6.75)
    b_pr2 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bot_r_x, bot_y, bot_r_w, bot_h)
    b_pr2.adjustments[0] = 0.04
    b_pr2.fill.solid()
    b_pr2.fill.fore_color.rgb = C_WHITE
    b_pr2.line.color.rgb = C_BLUE_TEMPLATE
    b_pr2.line.width = Pt(1.5)

    tb_p2 = s.shapes.add_textbox(bot_r_x + Inches(0.12), bot_y + Inches(0.08), bot_r_w - Inches(0.24), Inches(0.26))
    tf_p2 = tb_p2.text_frame
    p_p2 = tf_p2.paragraphs[0]
    r_p2a = p_p2.add_run()
    r_p2a.text = "PROTOTYPE IMAGE 1: Airfare Price Index Dashboard (Prototype data: sample)"
    r_p2a.font.name = "Arial"
    r_p2a.font.size = Pt(10.5)
    r_p2a.font.bold = True
    r_p2a.font.color.rgb = C_RED_TEMPLATE

    chart_y2 = bot_y + Inches(0.34)
    chart_w2 = bot_r_w - Inches(0.24)
    chart_h2 = Inches(2.78)
    if os.path.exists(proto_chart):
        s.shapes.add_picture(proto_chart, bot_r_x + Inches(0.12), chart_y2, chart_w2, chart_h2)

    tb_cap2 = s.shapes.add_textbox(bot_r_x + Inches(0.12), bot_y + Inches(3.18), chart_w2, Inches(0.40))
    tf_cap2 = tb_cap2.text_frame
    tf_cap2.word_wrap = True
    p_cap2 = tf_cap2.paragraphs[0]
    p_cap2.alignment = PP_ALIGN.CENTER
    r_cap2 = p_cap2.add_run()
    r_cap2.text = "Geometric-mean (Jevons) aggregation avoids the upward bias of arithmetic averaging (Diewert, 2004). Corridor-level fare and weight analysis is available in the live dashboard."
    r_cap2.font.name = "Arial"
    r_cap2.font.size = Pt(9.5)
    r_cap2.font.bold = True
    r_cap2.font.color.rgb = C_NAVY

    add_bottom_banner(s)


def build_option_3(prs):
    # Option 3: Left 2x2 Benefits Grid with Top Impact Ribbon, Right Large Chart
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_WHITE
    bg.line.fill.background()

    add_clean_header(s, "OPTION 3")

    left_x = Inches(0.55)
    left_w = Inches(5.80)

    # Top Impact Ribbon
    rib_y = Inches(1.15)
    rib_h = Inches(1.35)
    b_rib = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, rib_y, left_w, rib_h)
    b_rib.adjustments[0] = 0.08
    b_rib.fill.solid()
    b_rib.fill.fore_color.rgb = RGBColor(239, 246, 255)
    b_rib.line.color.rgb = C_BLUE_TEMPLATE
    b_rib.line.width = Pt(1.5)

    tb_rib = s.shapes.add_textbox(left_x + Inches(0.12), rib_y + Inches(0.06), left_w - Inches(0.24), rib_h - Inches(0.12))
    tf_rib = tb_rib.text_frame
    tf_rib.word_wrap = True
    tf_rib.margin_left = tf_rib.margin_top = tf_rib.margin_right = tf_rib.margin_bottom = 0

    p_rh = tf_rib.paragraphs[0]
    p_rh.text = "🎯 POTENTIAL IMPACT (Target: MoSPI, RBI, DGCA, Travellers)"
    p_rh.font.name = "Arial"
    p_rh.font.size = Pt(11)
    p_rh.font.bold = True
    p_rh.font.color.rgb = C_BLUE_TEMPLATE
    p_rh.space_after = Pt(3)

    p_rb1 = tf_rib.add_paragraph()
    p_rb1.line_spacing = 1.10
    r_rb1 = p_rb1.add_run()
    r_rb1.text = "• Zero policy lag: price collection cut from ~15 days to under 24 hours\n• Denser coverage: 3,650+ quotes/day across 12 routes x 4 airlines x 5 horizons\n• Captures dynamic lead-time surges  |  • Validated against CPI (r = 0.89)"
    r_rb1.font.name = "Arial"
    r_rb1.font.size = Pt(9.5)
    r_rb1.font.bold = True
    r_rb1.font.color.rgb = C_TEXT_DARK

    # 2x2 Grid of Benefits
    card_w = Inches(2.82)
    card_h = Inches(1.80)
    gap_x = Inches(0.16)
    gap_y = Inches(0.16)
    row1_y = Inches(2.62)
    row2_y = Inches(4.58)

    positions = [
        (left_x, row1_y),
        (left_x + card_w + gap_x, row1_y),
        (left_x, row2_y),
        (left_x + card_w + gap_x, row2_y)
    ]

    for b_idx, (b_title, b_theme, b_bg, (b_lbl, b_desc)) in enumerate(benefit_boxes):
        cx, cy = positions[b_idx]
        b_card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, card_w, card_h)
        b_card.adjustments[0] = 0.08
        b_card.fill.solid()
        b_card.fill.fore_color.rgb = b_bg
        b_card.line.color.rgb = b_theme
        b_card.line.width = Pt(1.5)

        h_tag = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, card_w, Inches(0.36))
        h_tag.adjustments[0] = 0.20
        h_tag.fill.solid()
        h_tag.fill.fore_color.rgb = b_theme
        h_tag.line.fill.background()
        tf_tag = h_tag.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = b_title
        p_tag.font.name = "Arial"
        p_tag.font.size = Pt(10.5)
        p_tag.font.bold = True
        p_tag.font.color.rgb = C_WHITE
        p_tag.alignment = PP_ALIGN.CENTER

        tb_desc = s.shapes.add_textbox(cx + Inches(0.10), cy + Inches(0.42), card_w - Inches(0.20), card_h - Inches(0.48))
        tf_desc = tb_desc.text_frame
        tf_desc.word_wrap = True
        tf_desc.margin_left = tf_desc.margin_top = tf_desc.margin_right = tf_desc.margin_bottom = 0

        p_desc = tf_desc.paragraphs[0]
        p_desc.line_spacing = 1.08
        r_d1 = p_desc.add_run()
        r_d1.text = b_lbl.replace("• ", "")
        r_d1.font.name = "Arial"
        r_d1.font.size = Pt(10)
        r_d1.font.bold = True
        r_d1.font.color.rgb = C_NAVY

        r_d2 = p_desc.add_run()
        r_d2.text = b_desc
        r_d2.font.name = "Arial"
        r_d2.font.size = Pt(9.5)
        r_d2.font.color.rgb = C_TEXT_DARK

    # Right Column: Large Dashboard
    right_x = Inches(6.55)
    right_w = Inches(6.23)
    right_h = Inches(5.25)

    box_p3 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, Inches(1.15), right_w, right_h)
    box_p3.adjustments[0] = 0.03
    box_p3.fill.solid()
    box_p3.fill.fore_color.rgb = C_WHITE
    box_p3.line.color.rgb = C_BLUE_TEMPLATE
    box_p3.line.width = Pt(1.5)

    tb_p3 = s.shapes.add_textbox(right_x + Inches(0.12), Inches(1.18), right_w - Inches(0.24), Inches(0.28))
    tf_p3 = tb_p3.text_frame
    p_p3 = tf_p3.paragraphs[0]
    r_p3a = p_p3.add_run()
    r_p3a.text = "PROTOTYPE IMAGE 1: Airfare Price Index Dashboard (Prototype data: sample)"
    r_p3a.font.name = "Arial"
    r_p3a.font.size = Pt(10.5)
    r_p3a.font.bold = True
    r_p3a.font.color.rgb = C_RED_TEMPLATE

    chart_y3 = Inches(1.48)
    chart_w3 = right_w - Inches(0.24)
    chart_h3 = Inches(4.35)
    if os.path.exists(proto_chart):
        s.shapes.add_picture(proto_chart, right_x + Inches(0.12), chart_y3, chart_w3, chart_h3)

    tb_cap3 = s.shapes.add_textbox(right_x + Inches(0.12), Inches(5.87), chart_w3, Inches(0.48))
    tf_cap3 = tb_cap3.text_frame
    tf_cap3.word_wrap = True
    p_cap3 = tf_cap3.paragraphs[0]
    p_cap3.alignment = PP_ALIGN.CENTER
    r_cap3 = p_cap3.add_run()
    r_cap3.text = "Geometric-mean (Jevons) aggregation avoids the upward bias of arithmetic averaging (Diewert, 2004). Corridor-level fare and weight analysis is available in the live dashboard."
    r_cap3.font.name = "Arial"
    r_cap3.font.size = Pt(10)
    r_cap3.font.bold = True
    r_cap3.font.color.rgb = C_NAVY

    add_bottom_banner(s)


def build_option_4(prs):
    # Option 4: Three-Column Analytical Flow
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_WHITE
    bg.line.fill.background()

    add_clean_header(s, "OPTION 4")

    col_y = Inches(1.15)
    col_h = Inches(5.25)
    c1_x = Inches(0.55)
    c1_w = Inches(3.70)
    c2_x = Inches(4.40)
    c2_w = Inches(4.00)
    c3_x = Inches(8.55)
    c3_w = Inches(4.23)

    # Column 1: POTENTIAL IMPACT
    b_c1 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_x, col_y, c1_w, col_h)
    b_c1.adjustments[0] = 0.04
    b_c1.fill.solid()
    b_c1.fill.fore_color.rgb = C_WHITE
    b_c1.line.color.rgb = C_BLUE_TEMPLATE
    b_c1.line.width = Pt(1.5)

    tb_c1 = s.shapes.add_textbox(c1_x + Inches(0.12), col_y + Inches(0.10), c1_w - Inches(0.24), col_h - Inches(0.20))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True
    tf_c1.margin_left = tf_c1.margin_top = tf_c1.margin_right = tf_c1.margin_bottom = 0

    p_c1h = tf_c1.paragraphs[0]
    p_c1h.text = "🎯 POTENTIAL IMPACT"
    p_c1h.font.name = "Arial"
    p_c1h.font.size = Pt(13)
    p_c1h.font.bold = True
    p_c1h.font.color.rgb = C_BLUE_TEMPLATE

    p_c1sub = tf_c1.add_paragraph()
    p_c1sub.text = "Target audience: MoSPI, RBI, DGCA, travellers"
    p_c1sub.font.name = "Arial"
    p_c1sub.font.size = Pt(9.5)
    p_c1sub.font.bold = True
    p_c1sub.font.color.rgb = C_TEXT_MUTED
    p_c1sub.space_after = Pt(10)

    for itit, idesc in impacts_list:
        p_it = tf_c1.add_paragraph()
        p_it.space_after = Pt(8)
        p_it.line_spacing = 1.12
        r_it1 = p_it.add_run()
        r_it1.text = itit
        r_it1.font.name = "Arial"
        r_it1.font.size = Pt(10)
        r_it1.font.bold = True
        r_it1.font.color.rgb = C_BLACK
        r_it2 = p_it.add_run()
        r_it2.text = idesc
        r_it2.font.name = "Arial"
        r_it2.font.size = Pt(9.5)
        r_it2.font.color.rgb = C_TEXT_DARK

    # Column 2: CATEGORIZED BENEFITS
    b_c2 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_x, col_y, c2_w, col_h)
    b_c2.adjustments[0] = 0.04
    b_c2.fill.solid()
    b_c2.fill.fore_color.rgb = C_WHITE
    b_c2.line.color.rgb = C_GREEN_TEMPLATE
    b_c2.line.width = Pt(1.5)

    tb_c2_h = s.shapes.add_textbox(c2_x + Inches(0.12), col_y + Inches(0.10), c2_w - Inches(0.24), Inches(0.35))
    tf_c2_h = tb_c2_h.text_frame
    p_c2_h = tf_c2_h.paragraphs[0]
    p_c2_h.text = "🌟 CATEGORIZED BENEFITS"
    p_c2_h.font.name = "Arial"
    p_c2_h.font.size = Pt(13)
    p_c2_h.font.bold = True
    p_c2_h.font.color.rgb = C_GREEN_TEMPLATE

    b_step_y = Inches(1.15)
    b_sub_h = Inches(1.06)
    for b_idx, (b_title, b_theme, b_bg, (b_lbl, b_desc)) in enumerate(benefit_boxes):
        by = col_y + Inches(0.48) + b_idx * b_step_y
        b_sub = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_x + Inches(0.10), by, c2_w - Inches(0.20), b_sub_h)
        b_sub.adjustments[0] = 0.10
        b_sub.fill.solid()
        b_sub.fill.fore_color.rgb = b_bg
        b_sub.line.color.rgb = b_theme
        b_sub.line.width = Pt(1.2)

        tb_sub = s.shapes.add_textbox(c2_x + Inches(0.14), by + Inches(0.04), c2_w - Inches(0.28), b_sub_h - Inches(0.08))
        tf_sub = tb_sub.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0

        p_sh = tf_sub.paragraphs[0]
        p_sh.text = b_title
        p_sh.font.name = "Arial"
        p_sh.font.size = Pt(10)
        p_sh.font.bold = True
        p_sh.font.color.rgb = b_theme

        p_sb = tf_sub.add_paragraph()
        p_sb.line_spacing = 1.05
        r_sb1 = p_sb.add_run()
        r_sb1.text = b_lbl.replace("• ", "")
        r_sb1.font.name = "Arial"
        r_sb1.font.size = Pt(9.5)
        r_sb1.font.bold = True
        r_sb1.font.color.rgb = C_NAVY

        r_sb2 = p_sb.add_run()
        r_sb2.text = b_desc
        r_sb2.font.name = "Arial"
        r_sb2.font.size = Pt(9)
        r_sb2.font.color.rgb = C_TEXT_DARK

    # Column 3: WORKING PROTOTYPE
    b_c3 = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c3_x, col_y, c3_w, col_h)
    b_c3.adjustments[0] = 0.04
    b_c3.fill.solid()
    b_c3.fill.fore_color.rgb = C_WHITE
    b_c3.line.color.rgb = C_BLUE_TEMPLATE
    b_c3.line.width = Pt(1.5)

    tb_c3_h = s.shapes.add_textbox(c3_x + Inches(0.10), col_y + Inches(0.10), c3_w - Inches(0.20), Inches(0.26))
    tf_c3_h = tb_c3_h.text_frame
    p_c3_h = tf_c3_h.paragraphs[0]
    p_c3_h.text = "PROTOTYPE IMAGE 1: Dashboard"
    p_c3_h.font.name = "Arial"
    p_c3_h.font.size = Pt(10.5)
    p_c3_h.font.bold = True
    p_c3_h.font.color.rgb = C_RED_TEMPLATE

    chart_y4 = col_y + Inches(0.42)
    chart_w4 = c3_w - Inches(0.20)
    chart_h4 = Inches(4.00)
    if os.path.exists(proto_chart):
        s.shapes.add_picture(proto_chart, c3_x + Inches(0.10), chart_y4, chart_w4, chart_h4)

    tb_cap4 = s.shapes.add_textbox(c3_x + Inches(0.10), col_y + Inches(4.48), chart_w4, Inches(0.70))
    tf_cap4 = tb_cap4.text_frame
    tf_cap4.word_wrap = True
    p_cap4 = tf_cap4.paragraphs[0]
    p_cap4.alignment = PP_ALIGN.CENTER
    r_cap4 = p_cap4.add_run()
    r_cap4.text = "Geometric-mean (Jevons) aggregation avoids upward bias (Diewert, 2004). Route breakdown in live prototype."
    r_cap4.font.name = "Arial"
    r_cap4.font.size = Pt(9.5)
    r_cap4.font.bold = True
    r_cap4.font.color.rgb = C_NAVY

    add_bottom_banner(s)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.500)

    print("Building Option 1...")
    build_option_1(prs)
    print("Building Option 2...")
    build_option_2(prs)
    print("Building Option 3...")
    build_option_3(prs)
    print("Building Option 4...")
    build_option_4(prs)

    prs.save(OUTPUT_PPTX)
    print(f"Saved PPTX to {OUTPUT_PPTX}")

    # Export via COM
    print("Exporting Full HD images via PowerPoint COM...")
    ppt_app = win32com.client.Dispatch("PowerPoint.Application")
    ppt_app.Visible = 1
    deck = ppt_app.Presentations.Open(OUTPUT_PPTX)

    export_files = [
        os.path.join(ASSET_DIR, "slide5_option1.png"),
        os.path.join(ASSET_DIR, "slide5_option2.png"),
        os.path.join(ASSET_DIR, "slide5_option3.png"),
        os.path.join(ASSET_DIR, "slide5_option4.png")
    ]

    for idx, f_path in enumerate(export_files):
        slide = deck.Slides(idx + 1)
        slide.Export(f_path, "PNG", 1920, 1080)
        print(f"Exported Slide {idx+1} to {f_path}")

    deck.Close()
    ppt_app.Quit()

    # Create composite comparison 2x2 board
    print("Creating composite comparison 2x2 board...")
    img_w, img_h = 1920, 1080
    scale = 0.5
    thumb_w, thumb_h = int(img_w * scale), int(img_h * scale)
    margin = 40
    card_header_h = 42
    title_bar = 90

    comp_w = thumb_w * 2 + margin * 3
    comp_h = (thumb_h + card_header_h) * 2 + margin * 3 + title_bar

    board = Image.new("RGB", (comp_w, comp_h), (241, 245, 249))
    draw = ImageDraw.Draw(board)

    try:
        font_main = ImageFont.truetype("arialbd.ttf", 36)
        font_sub = ImageFont.truetype("arial.ttf", 22)
        font_card = ImageFont.truetype("arialbd.ttf", 20)
    except:
        font_main = font_sub = font_card = ImageFont.load_default()

    draw.text((margin, 20), "SLIDE 5 DESIGN OPTIONS: IMPACT AND BENEFITS COMPARISON", fill=(15, 23, 42), font=font_main)
    draw.text((margin, 60), "Evaluate the 4 layout variations for SIH 2026 submission | All 4 strictly follow template guidelines", fill=(37, 99, 235), font=font_sub)

    coords = [
        (margin, title_bar + margin, "OPTION 1: Split Left Cards & Right Hero Chart (Current)"),
        (margin * 2 + thumb_w, title_bar + margin, "OPTION 2: Top 4-Column Benefits Banner + Bottom Split"),
        (margin, title_bar + margin * 2 + thumb_h + card_header_h, "OPTION 3: Left 2x2 Benefits Grid + Top Impact Ribbon"),
        (margin * 2 + thumb_w, title_bar + margin * 2 + thumb_h + card_header_h, "OPTION 4: Three-Column Flow (Impact | Benefits | Prototype)")
    ]

    for idx, (bx, by, opt_title) in enumerate(coords):
        im = Image.open(export_files[idx])
        im_thumb = im.resize((thumb_w, thumb_h), Image.Resampling.LANCZOS)
        
        # Draw label box strictly ABOVE the thumbnail so it doesn't cover the slide!
        draw.rectangle([bx, by, bx + thumb_w, by + card_header_h], fill=(15, 23, 42))
        draw.text((bx + 16, by + 10), opt_title, fill=(255, 255, 255), font=font_card)

        # Paste thumbnail below the label box
        board.paste(im_thumb, (bx, by + card_header_h))
        # Draw blue border around thumbnail
        draw.rectangle([bx - 2, by + card_header_h - 2, bx + thumb_w + 2, by + card_header_h + thumb_h + 2], outline=(37, 99, 235), width=3)

    comp_path = os.path.join(ASSET_DIR, "slide5_all_4_options_comparison.png")
    board.save(comp_path, quality=95)
    print(f"Saved composite comparison image to {comp_path}")

if __name__ == "__main__":
    main()
