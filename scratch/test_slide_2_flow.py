import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def test_slide_2():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    C_WHITE = RGBColor(255, 255, 255)
    C_BLACK = RGBColor(0, 0, 0)
    C_BLUE_TEMPLATE = RGBColor(30, 64, 175)
    C_TEXT_DARK = RGBColor(15, 23, 42)

    base_dir = r"C:\Users\oshsh\.gemini\antigravity\scratch\sih26056_airfare_cpi"
    pres_dir = os.path.join(base_dir, "public", "presentation")
    sih_top_logo = os.path.join(pres_dir, "page_1_img_2.png")

    s2 = prs.slides.add_slide(blank_layout)
    bg2 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = C_WHITE
    bg2.line.fill.background()

    # Top Bar: Roorkies Pill, Title, SIH Logo
    pill = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(0.22), Inches(1.8), Inches(0.55))
    pill.adjustments[0] = 0.5
    pill.fill.solid()
    pill.fill.fore_color.rgb = C_WHITE
    pill.line.color.rgb = RGBColor(100, 116, 139)
    pill.line.width = Pt(1.5)
    tf_p = pill.text_frame
    tf_p.word_wrap = True
    p_p = tf_p.paragraphs[0]
    p_p.text = "Roorkies"
    p_p.font.name = "Arial"
    p_p.font.size = Pt(13)
    p_p.font.bold = True
    p_p.font.color.rgb = C_BLACK
    p_p.alignment = PP_ALIGN.CENTER

    tb_t = s2.shapes.add_textbox(Inches(2.5), Inches(0.12), Inches(8.3), Inches(0.95))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    p_t1 = tf_t.paragraphs[0]
    p_t1.text = "VayuSuchak: Real-Time Airfare Price Index"
    p_t1.font.name = "Arial"
    p_t1.font.size = Pt(21)
    p_t1.font.bold = True
    p_t1.font.color.rgb = C_BLACK
    p_t1.alignment = PP_ALIGN.CENTER

    p_t2 = tf_t.add_paragraph()
    p_t2.text = "Augmenting CPI Transport Inflation"
    p_t2.font.name = "Arial"
    p_t2.font.size = Pt(15)
    p_t2.font.bold = True
    p_t2.font.color.rgb = C_BLACK
    p_t2.alignment = PP_ALIGN.CENTER

    if os.path.exists(sih_top_logo):
        s2.shapes.add_picture(sih_top_logo, Inches(11.1), Inches(0.12), Inches(1.8), Inches(0.85))

    # Subtitle
    tb_s2_sub = s2.shapes.add_textbox(Inches(0.55), Inches(1.08), Inches(12.2), Inches(0.35))
    tf_s2_sub = tb_s2_sub.text_frame
    p_s2_sub = tf_s2_sub.paragraphs[0]
    p_s2_sub.text = "❖ Proposed Solution: Problem vs. Solution Architecture"
    p_s2_sub.font.name = "Arial"
    p_s2_sub.font.size = Pt(15)
    p_s2_sub.font.bold = True
    p_s2_sub.font.color.rgb = C_BLUE_TEMPLATE

    # Top Row: Problem & Solution
    top_y = Inches(1.48)
    top_h = Inches(2.62)
    col_w = Inches(5.95)

    # Problem Card
    box_prob = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), top_y, col_w, top_h)
    box_prob.adjustments[0] = 0.04
    box_prob.fill.solid()
    box_prob.fill.fore_color.rgb = RGBColor(254, 242, 242)
    box_prob.line.color.rgb = RGBColor(220, 38, 38)
    box_prob.line.width = Pt(1.5)

    hdr_prob = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), top_y, col_w, Inches(0.42))
    hdr_prob.adjustments[0] = 0.2
    hdr_prob.fill.solid()
    hdr_prob.fill.fore_color.rgb = RGBColor(220, 38, 38)
    hdr_prob.line.fill.background()
    p_ph = hdr_prob.text_frame.paragraphs[0]
    p_ph.text = "⚠️ THE PROBLEM (Current MoSPI Manual Survey)"
    p_ph.font.name = "Arial"
    p_ph.font.size = Pt(11.5)
    p_ph.font.bold = True
    p_ph.font.color.rgb = C_WHITE
    p_ph.alignment = PP_ALIGN.CENTER

    tb_prob = s2.shapes.add_textbox(Inches(0.75), top_y + Inches(0.48), col_w - Inches(0.40), top_h - Inches(0.52))
    tf_prob = tb_prob.text_frame
    tf_prob.word_wrap = True

    prob_points = [
        ("15-Day Information Lag: ", "Manual surveyor visits delay CPI reporting by 2 weeks, missing high-frequency price volatility."),
        ("Static Single Snapshot: ", "Only 1 quote collected per route/month, failing to capture 10,000+ daily dynamic algorithmic fares."),
        ("Dutot Upward Bias: ", "Arithmetic mean formula overstates airfare transport inflation by 20–30 bps (substitution bias)."),
        ("High Operational Cost: ", "Multi-crore physical surveyor logistics across airports with zero cryptographic audit trail.")
    ]
    for idx, (head, desc) in enumerate(prob_points):
        p = tf_prob.paragraphs[0] if idx == 0 else tf_prob.add_paragraph()
        p.space_after = Pt(3.5)
        p.line_spacing = 1.15
        r1 = p.add_run()
        r1.text = "• " + head
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(185, 28, 28)
        r2 = p.add_run()
        r2.text = desc
        r2.font.name = "Arial"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = C_TEXT_DARK

    # Solution Card
    r_x = Inches(6.83)
    box_sol = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, r_x, top_y, col_w, top_h)
    box_sol.adjustments[0] = 0.04
    box_sol.fill.solid()
    box_sol.fill.fore_color.rgb = RGBColor(240, 253, 244)
    box_sol.line.color.rgb = RGBColor(22, 163, 74)
    box_sol.line.width = Pt(1.5)

    hdr_sol = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, r_x, top_y, col_w, Inches(0.42))
    hdr_sol.adjustments[0] = 0.2
    hdr_sol.fill.solid()
    hdr_sol.fill.fore_color.rgb = RGBColor(22, 163, 74)
    hdr_sol.line.fill.background()
    p_sh = hdr_sol.text_frame.paragraphs[0]
    p_sh.text = "💡 HOW OUR APP SOLVES IT (VayuSuchak Engine)"
    p_sh.font.name = "Arial"
    p_sh.font.size = Pt(11.5)
    p_sh.font.bold = True
    p_sh.font.color.rgb = C_WHITE
    p_sh.alignment = PP_ALIGN.CENTER

    tb_sol = s2.shapes.add_textbox(r_x + Inches(0.20), top_y + Inches(0.48), col_w - Inches(0.40), top_h - Inches(0.52))
    tf_sol = tb_sol.text_frame
    tf_sol.word_wrap = True

    sol_points = [
        ("Real-Time Extraction: ", "Playwright headless bots harvest 10,000+ live fares daily with sub-24h latency at 02:00 AM IST."),
        ("T+1..T+45 Advance Horizons: ", "Captures urgent vs saver fares weighted dynamically by official DGCA passenger volumes."),
        ("UN/ILO Jevons Index: ", "Geometric mean formulation mathematically eliminates Dutot upward substitution distortion."),
        ("SHA-256 Audit Ledger: ", "Immutable cryptographic fingerprints ensure sovereign-grade evidentiary auditability for MoSPI.")
    ]
    for idx, (head, desc) in enumerate(sol_points):
        p = tf_sol.paragraphs[0] if idx == 0 else tf_sol.add_paragraph()
        p.space_after = Pt(3.5)
        p.line_spacing = 1.15
        r1 = p.add_run()
        r1.text = "✓ " + head
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(21, 128, 61)
        r2 = p.add_run()
        r2.text = desc
        r2.font.name = "Arial"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = C_TEXT_DARK

    # --- CONNECTED 5-STEP HORIZONTAL PROCESS WORKFLOW ---
    flow_y = Inches(4.25)
    flow_h = Inches(1.95)
    
    total_w = Inches(12.23)
    arrow_w = Inches(0.20)
    gap_arrow = Inches(0.04)
    total_arrows_w = 4 * (arrow_w + 2 * gap_arrow)
    card_w = (total_w - total_arrows_w) / 5

    steps = [
        (
            "STEP 1",
            "Dynamic Yield Tracking",
            "Captures intraday surges across T+1..T+45 advance booking windows.",
            RGBColor(30, 64, 175), # Blue
            RGBColor(239, 246, 255)
        ),
        (
            "STEP 2",
            "Dynamic IQR Scrubber",
            "Isolates base fares and purges phantom prices and cache glitches.",
            RGBColor(124, 58, 237), # Purple
            RGBColor(245, 243, 255)
        ),
        (
            "STEP 3",
            "UN/ILO Jevons Index",
            "Geometric mean calculation that eliminates upward substitution bias.",
            RGBColor(2, 132, 199), # Sky
            RGBColor(240, 249, 255)
        ),
        (
            "STEP 4",
            "Sovereign Audit Vault",
            "SHA-256 batch Merkle tree proofs for judicial and policy scrutiny.",
            RGBColor(16, 185, 129), # Emerald
            RGBColor(236, 253, 245)
        ),
        (
            "STEP 5",
            "Real-Time API Delivery",
            "Sub-10ms REST API and live eSankhyiki CSV feed for MoSPI and RBI.",
            RGBColor(217, 119, 6), # Amber
            RGBColor(254, 252, 232)
        )
    ]

    for idx, (s_num, s_title, s_desc, border_c, bg_c) in enumerate(steps):
        cx = Inches(0.55) + idx * (card_w + arrow_w + 2 * gap_arrow)
        
        # Step Card
        card_shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, flow_y, card_w, flow_h)
        card_shape.adjustments[0] = 0.08
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = bg_c
        card_shape.line.color.rgb = border_c
        card_shape.line.width = Pt(1.4)

        # Top Badge Pill
        badge = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Inches(0.12), flow_y + Inches(0.10), Inches(0.85), Inches(0.26))
        badge.adjustments[0] = 0.5
        badge.fill.solid()
        badge.fill.fore_color.rgb = border_c
        badge.line.fill.background()
        p_b = badge.text_frame.paragraphs[0]
        p_b.text = s_num
        p_b.font.name = "Arial"
        p_b.font.size = Pt(8.5)
        p_b.font.bold = True
        p_b.font.color.rgb = C_WHITE
        p_b.alignment = PP_ALIGN.CENTER

        # Content Textbox
        tb_c = s2.shapes.add_textbox(cx + Inches(0.12), flow_y + Inches(0.40), card_w - Inches(0.24), flow_h - Inches(0.46))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0

        p_ct = tf_c.paragraphs[0]
        p_ct.text = s_title
        p_ct.font.name = "Arial"
        p_ct.font.size = Pt(10.5)
        p_ct.font.bold = True
        p_ct.font.color.rgb = border_c
        p_ct.space_after = Pt(4)

        p_cd = tf_c.add_paragraph()
        p_cd.text = s_desc
        p_cd.font.name = "Arial"
        p_cd.font.size = Pt(8.8)
        p_cd.font.color.rgb = C_TEXT_DARK
        p_cd.line_spacing = 1.15

        # Connecting Arrow between cards (for first 4 cards)
        if idx < 4:
            ax = cx + card_w + gap_arrow
            ay = flow_y + flow_h / 2 - Inches(0.14)
            arrow = s2.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax, ay, arrow_w, Inches(0.28))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor(148, 163, 184) # slate-400
            arrow.line.fill.background()

    # --- BOTTOM ROW: PROCESS PIPELINE BANNER ---
    bot_y = Inches(6.35)
    bot_h = Inches(0.50)
    box_bottom = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), bot_y, Inches(12.23), bot_h)
    box_bottom.adjustments[0] = 0.2
    box_bottom.fill.solid()
    box_bottom.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box_bottom.line.color.rgb = C_BLUE_TEMPLATE
    box_bottom.line.width = Pt(1.5)

    tb_bb = s2.shapes.add_textbox(Inches(0.65), bot_y + Inches(0.06), Inches(12.03), Inches(0.38))
    tf_bb = tb_bb.text_frame
    p_bb = tf_bb.paragraphs[0]
    p_bb.text = "DATA HARVESTING  ➔  IQR OUTLIER TRUNCATION  ➔  JEVONS GEOMETRIC INDEX  ➔  SHA-256 PROVENANCE  ➔  MoSPI & RBI CPI INTEGRATION"
    p_bb.font.name = "Arial"
    p_bb.font.size = Pt(11)
    p_bb.font.bold = True
    p_bb.font.color.rgb = C_BLACK
    p_bb.alignment = PP_ALIGN.CENTER

    out_pptx = os.path.join(base_dir, "scratch", "test_slide_2_flow.pptx")
    prs.save(out_pptx)
    print("Saved test pptx to", out_pptx)

    # Export via COM
    import win32com.client
    ppt = win32com.client.Dispatch("PowerPoint.Application")
    pres_com = ppt.Presentations.Open(os.path.abspath(out_pptx), False, False, False)
    out_png = os.path.join(base_dir, "scratch", "test_slide_2_flow.png")
    pres_com.Slides(1).Export(out_png, "PNG", 1920, 1080)
    pres_com.Close()
    ppt.Quit()
    print("Exported test png to", out_png)

if __name__ == "__main__":
    test_slide_2()
